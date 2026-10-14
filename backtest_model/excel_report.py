# -*- coding: utf-8 -*-
"""
==============================================================================
Excel 导出 —— 多 Sheet + 表头样式 + 列宽 + 涨跌着色 + 文本格式
==============================================================================

与迁移来源 `three_factor_best_export.py` / `three_factor_best_backtest.py`
的样式约定保持一致（表头 `366092` 底 + 白色加粗字；正绿负红）。
"""

import pandas as pd

from . import settings as S

_RED = 'C00000'
_GREEN = '006100'
_NEG_FILL = 'FCE4E4'
_POS_FILL = 'E2EFDA'


def _style_sheet(ws, df, pct_cols=(), text_cols=(), color_cols=(), freeze=None,
                 min_width=9, width_cap=40):
    """统一表头样式、列宽、数字格式与涨跌着色。

    Args:
        ws: openpyxl worksheet
        df: 对应 DataFrame
        pct_cols: 需要固定为 '0.00' 数字格式的百分比列
        text_cols: 需要强制文本格式（保留前导零，如股票代码）的列
        color_cols: 需要按正负着色的列
        freeze: 冻结窗格（如 'A2'）
        min_width / width_cap: 列宽上下限
    """
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color=S.HEADER_FILL,
                              end_color=S.HEADER_FILL, fill_type='solid')
    header_font = Font(color=S.HEADER_FONT, bold=True)
    header_align = Alignment(horizontal='center', vertical='center',
                             wrap_text=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # 列宽（中文按 2 字符宽计）
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        head_len = sum(2 if ord(ch) > 255 else 1 for ch in str(col))
        if len(df):
            # 注意：pandas 3.x（infer_string）下 astype(str) 会把缺失值保留为
            # NaN（而非 'nan'），故此处显式判空，避免迭代到非字符串对象。
            body_len = df[col].head(200).map(
                lambda v: sum(2 if ord(ch) > 255 else 1
                              for ch in ('' if pd.isna(v) else str(v)))).max()
        else:
            body_len = 0
        ws.column_dimensions[letter].width = max(
            min_width, min(max(head_len, body_len) + 2, width_cap))

    red = Font(color=_RED)
    green = Font(color=_GREEN)
    neg_fill = PatternFill(start_color=_NEG_FILL, end_color=_NEG_FILL,
                           fill_type='solid')
    pos_fill = PatternFill(start_color=_POS_FILL, end_color=_POS_FILL,
                           fill_type='solid')

    # 涨跌着色
    for name in color_cols:
        if name not in df.columns or not len(df):
            continue
        ci = df.columns.get_loc(name) + 1
        for ri in range(2, len(df) + 2):
            cell = ws.cell(row=ri, column=ci)
            if isinstance(cell.value, (int, float)):
                if cell.value > 0:
                    cell.font, cell.fill = green, pos_fill
                elif cell.value < 0:
                    cell.font, cell.fill = red, neg_fill

    # 百分比格式
    for name in pct_cols:
        if name not in df.columns:
            continue
        ci = df.columns.get_loc(name) + 1
        for ri in range(2, len(df) + 2):
            ws.cell(row=ri, column=ci).number_format = '0.00'

    # 文本格式（如股票代码保留前导零）
    for name in text_cols:
        if name not in df.columns:
            continue
        ci = df.columns.get_loc(name) + 1
        for ri in range(2, len(df) + 2):
            cell = ws.cell(row=ri, column=ci)
            if cell.value is not None:
                cell.number_format = '@'

    if freeze:
        ws.freeze_panes = freeze


def export_excel(sheets, out_path):
    """把 [(名称, DataFrame, 选项)] 列表写入多 Sheet Excel。

    选项键：`pct` / `text` / `color` / `freeze` / `cap`。
    """
    print('\n📊 正在写入 Excel ...')
    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        for name, d, opt in sheets:
            d.to_excel(writer, sheet_name=name, index=False)
            _style_sheet(
                writer.sheets[name], d,
                pct_cols=opt.get('pct', ()),
                text_cols=opt.get('text', ()),
                color_cols=opt.get('color', ()),
                freeze=opt.get('freeze'),
                width_cap=opt.get('cap', 40),
            )
            print(f'   · {name:<10} {len(d):>6} 行 × {len(d.columns)} 列')
    return out_path
